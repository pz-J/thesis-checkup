"""生成一份格式全部合规的“满分论文”，用于回归测试与误报校准。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from make_sample import make_figure

OUT = Path(__file__).resolve().parent.parent / "public" / "sample-clean.docx"
IMG = Path(__file__).resolve().parent / "_clean-figure.png"


def set_style_font(style, east_asia: str, size_pt: float, bold: bool = False, ascii_font: str = "Times New Roman") -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def apply_run_font(run, east="宋体", size=12, bold=False, ascii_font="Times New Roman"):
    run.bold = bold
    run.font.name = ascii_font
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east)


def body_para(doc, text: str, *, indent=True, east="宋体", size=12, ascii_font="Times New Roman"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = Pt(20)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = Pt(24)
    run = p.add_run(text)
    apply_run_font(run, east=east, size=size, ascii_font=ascii_font)
    return p


def title_para(doc, text: str, *, east="黑体", size=16, bold=False, ascii_font="Times New Roman", page_break=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    apply_run_font(run, east=east, size=size, bold=bold, ascii_font=ascii_font)
    return p


def heading(doc, text: str, level: int, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text)
    return p


def add_field(paragraph, instruction: str, placeholder: str = "1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(text)
    run._r.append(end)
    apply_run_font(run, east="宋体", size=10.5)


def add_page_number_footer(doc):
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ", "1")


def main() -> None:
    make_figure(IMG)

    doc = Document()
    doc.core_properties.title = "面向高校论文的格式自动体检方法"
    doc.core_properties.author = "满分示例"

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    set_style_font(normal, "宋体", 12)
    normal.paragraph_format.line_spacing = Pt(20)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    set_style_font(doc.styles["Heading 1"], "黑体", 14)
    set_style_font(doc.styles["Heading 2"], "黑体", 12)
    set_style_font(doc.styles["Heading 3"], "楷体", 12)

    # 封面
    for _ in range(4):
        doc.add_paragraph()
    title_para(doc, "天津大学本科毕业设计（论文）", east="黑体", size=22, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    for line in [
        "题目：面向高校论文的格式自动体检方法",
        "学院：智能与计算学部",
        "专业：计算机科学与技术",
        "学生姓名：满分示例",
        "学号：3026XXXXXX",
        "指导教师：示例教师",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        apply_run_font(run, east="宋体", size=16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_run_font(p.add_run("2026 年 6 月"), east="宋体", size=16)

    # 中文摘要
    title_para(doc, "摘  要", east="黑体", size=16, page_break=True)
    body_para(
        doc,
        "论文格式是毕业设计提交前的最后一道硬门槛。本文设计并实现了一个完全在浏览器本地运行的论文格式体检工具，"
        "它直接解析 Word 文档的 OOXML 结构，按学校规则包检查页面、正文、标题、图表题注与参考文献，"
        "并生成可视化报告与可自动修复的文档副本。工具不上传论文，不依赖账号，适合在提交前完成一次快速自查。",
        indent=False,
    )
    kw = doc.add_paragraph()
    kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    kw.paragraph_format.line_spacing = Pt(20)
    apply_run_font(kw.add_run("关键词："), east="黑体", size=12)
    apply_run_font(kw.add_run("论文格式；Word；OOXML；自动检查"), east="宋体", size=12)

    # 英文摘要
    title_para(doc, "Abstract", east="Times New Roman", size=16, bold=True, ascii_font="Times New Roman")
    body_para(
        doc,
        "Formatting is the last hard checkpoint before a thesis is submitted. This paper presents a local-first "
        "thesis format checker that parses OOXML in the browser, evaluates school-specific rules, and produces "
        "a visual report together with a repaired copy of the document.",
        indent=False,
        east="Times New Roman",
        ascii_font="Times New Roman",
    )
    kw2 = doc.add_paragraph()
    kw2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    kw2.paragraph_format.line_spacing = Pt(20)
    apply_run_font(kw2.add_run("Key words: "), east="Times New Roman", size=12, ascii_font="Times New Roman")
    apply_run_font(kw2.add_run("thesis format; Word; OOXML; automated check"), east="Times New Roman", size=12, ascii_font="Times New Roman")

    # 目录
    title_para(doc, "目  录", east="黑体", size=16)
    toc_p = doc.add_paragraph()
    toc_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_field(toc_p, ' TOC \\o "1-3" \\h \\z \\u ', "右键更新目录")
    doc.add_page_break()

    # 正文
    heading(doc, "1 绪论", 1)
    heading(doc, "1.1 研究背景", 2)
    body_para(
        doc,
        "每年毕业季，格式问题都会消耗学生与指导教师大量时间。页边距、字体、行距、标题层级、图表编号和参考文献著录，"
        "任何一项不符合模板，都可能在形式审查阶段被退回。现有工具要么依赖上传文档到云端，要么只覆盖单一学校规则，"
        "学生很难在提交前快速完成一次全面自查。",
    )
    body_para(
        doc,
        "更麻烦的是，Word 的格式信息具有明显的继承性。一个段落看起来是宋体小四，实际可能来自主题字体、样式继承或"
        "局部直接格式的叠加；肉眼很难判断问题究竟出在哪一层。只有直接读取 OOXML，才能把「看起来对」和「实际对」区分开。",
    )
    heading(doc, "1.1.1 研究意义", 3)
    body_para(
        doc,
        "本文的意义在于把论文格式检查从「人工逐项比对」变成「一次结构化体检」。工具不依赖云端，不上传论文，"
        "所有解析与修复都在浏览器内完成，既保护隐私，也降低了使用门槛。",
    )

    # 图与题注
    doc.add_picture(str(IMG), width=Cm(12))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para(doc, "图1 系统总体架构", east="宋体", size=10.5)

    # 表与题注
    title_para(doc, "表1 实验数据对比", east="宋体", size=10.5)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ("项目", "传统人工检查", "本文工具"),
        ("平均耗时", "35 分钟", "40 秒"),
        ("漏检项", "4.2 项", "0.6 项"),
        ("隐私风险", "需上传云端", "本地完成"),
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for run in cell.paragraphs[0].runs:
                apply_run_font(run, east="宋体", size=10.5)

    heading(doc, "2 系统设计", 1, page_break=True)
    body_para(
        doc,
        "系统由解析层、规则层、修复层和展示层组成。解析层读取文档的样式表与正文 XML，重建每个段落的有效格式；"
        "规则层把学校模板表达为可配置的 JSON 规则包；修复层只修改确定性强的问题，并保留其余内容不变。",
    )
    heading(doc, "2.1 模块划分", 2)
    body_para(
        doc,
        "解析层负责读取 OOXML；规则层负责把学校规范转成可执行检查；修复层负责生成新的文档副本；展示层负责把结果呈现为可读的报告。",
    )
    heading(doc, "3 结论", 1, page_break=True)
    body_para(
        doc,
        "论文格式体检把隐性的排版知识变成可执行的规则。下一步可以扩展学校规则包、增加更多自动修复项，"
        "并把检查能力接入写作流程，在提交前自动提醒。",
    )

    # 参考文献（标题顶格，不居中）
    ref_title = doc.add_paragraph()
    ref_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ref_title.paragraph_format.page_break_before = True
    ref_title.paragraph_format.space_before = Pt(0)
    ref_title.paragraph_format.space_after = Pt(0)
    apply_run_font(ref_title.add_run("参考文献"), east="黑体", size=14)

    for ref in [
        "[1] 张三. 论文格式自动检查研究[J]. 天津大学学报, 2020, 53(3): 1-7.",
        "[2] 李四. 深度学习[M]. 北京: 科学出版社, 2019: 58-62.",
        "[3] 王五. 毕业论文排版方法[D]. 天津: 天津大学, 2021.",
    ]:
        rp = doc.add_paragraph()
        rp.paragraph_format.line_spacing = Pt(20)
        apply_run_font(rp.add_run(ref), east="宋体", size=10.5)

    # 谢辞
    title_para(doc, "谢  辞", east="黑体", size=14, page_break=True)
    body_para(doc, "感谢导师在论文写作与格式规范方面的耐心指导，也感谢同学在测试阶段提供的真实论文样本。", indent=False)

    # 附录
    title_para(doc, "附  录", east="黑体", size=14, page_break=True)
    body_para(doc, "附录 A：规则包字段说明。规则包采用 JSON 描述页面、正文、标题、题注与参考文献的期望值。", indent=False)

    add_page_number_footer(doc)
    doc.save(OUT)
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
