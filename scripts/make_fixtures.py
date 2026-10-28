"""生成真实结构测试样本：已有页脚、多节文档、现成 TOC 域。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)


def add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(text)
    run._r.append(end)
    run.font.size = Pt(10.5)


def setup_a4(doc: Document) -> None:
    for sec in doc.sections:
        sec.page_width = Cm(21)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)


def existing_footer() -> None:
    doc = Document()
    setup_a4(doc)
    d = doc.add_paragraph("这是正文段落，用来验证已有页脚时不会覆盖原有内容。")
    for run in d.runs:
        run.font.size = Pt(12)
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "天津大学"
    footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(OUT / "existing-footer.docx")


def multi_section() -> None:
    doc = Document()
    setup_a4(doc)
    first = doc.sections[0]
    first.footer.is_linked_to_previous = False
    add_field(first.footer.paragraphs[0], " PAGE ", "1")
    doc.add_paragraph("第一节正文，已有自动页码。")
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = "第二节页脚"
    doc.add_paragraph("第二节正文，没有自动页码。")
    setup_a4(doc)
    doc.save(OUT / "multi-section.docx")


def toc_field() -> None:
    doc = Document()
    setup_a4(doc)
    doc.add_paragraph("目录")
    p = doc.add_paragraph()
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ', "右键更新目录")
    doc.add_heading("1 绪论", 1)
    doc.add_paragraph("正文内容。")
    doc.add_heading("2 结论", 1)
    doc.add_paragraph("结论内容。")
    doc.save(OUT / "toc-field.docx")


def main() -> None:
    existing_footer()
    multi_section()
    toc_field()
    print(f"written fixtures to {OUT}")


if __name__ == "__main__":
    main()
