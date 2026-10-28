"""生成一份故意带格式问题的示例论文，用于体检演示与回归测试。"""
from __future__ import annotations

import struct
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent.parent / "public" / "sample-thesis.docx"
IMG = Path(__file__).resolve().parent / "_sample-figure.png"


def png_chunk(tag: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)


def make_figure(path: Path, w: int = 640, h: int = 320) -> None:
    rows = bytearray()
    for y in range(h):
        rows.append(0)
        for x in range(w):
            bg = (246, 244, 239)
            r, g, b = bg
            # 网格
            if x % 64 == 0 or y % 64 == 0:
                r, g, b = (232, 229, 221)
            # 柱状图
            bars = [
                (90, 180, 110, 70),
                (190, 220, 150, 90),
                (290, 150, 95, 55),
                (390, 240, 210, 120),
                (490, 200, 120, 80),
            ]
            for bx, bh, col_r, col_g in bars:
                if bx <= x <= bx + 46 and h - 40 - bh <= y <= h - 40:
                    r, g, b = (col_r, col_g, 120)
            # 坐标轴
            if 40 <= x <= w - 30 and abs(y - (h - 40)) <= 1:
                r, g, b = (90, 90, 86)
            rows += bytes((r, g, b))
    data = b"\x89PNG\r\n\x1a\n"
    data += png_chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
    data += png_chunk(b"IDAT", zlib.compress(bytes(rows), 9))
    data += png_chunk(b"IEND", b"")
    path.write_bytes(data)


def set_style_font(style, east_asia: str, size_pt: float, bold: bool = False, ascii_font: str = "Times New Roman") -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def para(doc, text: str, *, east="宋体", size=12, bold=False, align=None, ascii_font="Times New Roman", space_after=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = ascii_font
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east)
    return p


def heading(doc, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_page_break()


def add_page_number_field_footer(doc):
    """示例论文故意不调用：保留“没有页码”的问题。"""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def main() -> None:
    make_figure(IMG)

    doc = Document()
    doc.core_properties.title = "面向高校论文的格式自动体检方法"
    doc.core_properties.author = "示例学生"

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)  # 故意错误：天大规则要求 2.5
    section.right_margin = Cm(2.0)  # 故意错误

    normal = doc.styles["Normal"]
    set_style_font(normal, "宋体", 10.5, bold=False)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 标题样式故意混入错误字体
    set_style_font(doc.styles["Heading 1"], "宋体", 14, bold=False)
    set_style_font(doc.styles["Heading 2"], "黑体", 12, bold=False)
    set_style_font(doc.styles["Heading 3"], "宋体", 12, bold=False)

    # 封面
    for _ in range(4):
        doc.add_paragraph()
    para(doc, "天津大学本科毕业设计（论文）", east="黑体", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        doc.add_paragraph()
    para(doc, "题目：面向高校论文的格式自动体检方法", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "学院：智能与计算学部", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "专业：计算机科学与技术", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "学生姓名：示例学生", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "学号：3026XXXXXX", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "指导教师：示例教师", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026 年 6 月", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # 中文摘要（标题字体故意错误）
    para(doc, "摘  要", east="宋体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        doc,
        "论文格式是毕业设计提交前的最后一道硬门槛。本文设计并实现了一个完全在浏览器本地运行的论文格式体检工具，"
        "它直接解析 Word 文档的 OOXML 结构，按学校规则包检查页面、正文、标题、图表题注与参考文献，"
        "并生成可视化报告与可自动修复的文档副本。",
        east="仿宋",
        size=10.5,
    )
    para(doc, "关键词：论文格式；Word；OOXML；自动检查", east="宋体", size=12)

    # 英文摘要
    para(doc, "Abstract", east="Times New Roman", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ascii_font="Times New Roman")
    para(
        doc,
        "Formatting is the last hard checkpoint before a thesis is submitted. This paper presents a local-first "
        "thesis format checker that parses OOXML in the browser, evaluates school-specific rules, and produces "
        "a visual report together with a repaired copy of the document.",
        east="Times New Roman",
        size=12,
        ascii_font="Times New Roman",
    )
    para(doc, "Key words: thesis format; Word; OOXML; automated check", east="Times New Roman", size=12, ascii_font="Times New Roman")

    # 目录（静态文本，无 TOC 域）
    para(doc, "目  录", east="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in [
        "1 绪论 ……………………………………………………………………… 1",
        "1.1 研究背景 ……………………………………………………………… 1",
        "1.1.1 研究意义 …………………………………………………………… 2",
        "2 系统设计 ………………………………………………………………… 3",
        "3 结论 ……………………………………………………………………… 6",
        "谢辞 ………………………………………………………………………… 7",
        "参考文献 …………………………………………………………………… 8",
        "附录 ………………………………………………………………………… 9",
    ]:
        para(doc, line, east="宋体", size=12)
    page_break(doc)

    # 正文
    heading(doc, "1 绪论", 1)
    heading(doc, "1.1 研究背景", 2)
    para(
        doc,
        "每年毕业季，格式问题都会消耗学生与指导教师大量时间。页边距、字体、行距、标题层级、图表编号和参考文献著录，"
        "任何一项不符合模板，都可能在形式审查阶段被退回。已有研究[9]尝试过自动检查，但现有工具要么依赖上传文档到云端，"
        "要么只覆盖单一学校规则，学生很难在提交前快速完成一次全面自查。",
        east="仿宋",
        size=10.5,
    )
    para(
        doc,
        "　　更麻烦的是，Word 的格式信息具有明显的继承性。一个段落看起来是宋体小四，实际可能来自主题字体、样式继承或"
        "局部直接格式的叠加；肉眼很难判断问题究竟出在哪一层。只有直接读取 OOXML，才能把「看起来对」和「实际对」区分开。",
        east="仿宋",
        size=10.5,
    )
    heading(doc, "1.1.1 研究意义", 3)
    para(
        doc,
        "本文的意义在于把论文格式检查从「人工逐项比对」变成「一次结构化体检」。工具不依赖云端，不上传论文，"
        "所有解析与修复都在浏览器内完成，既保护隐私，也降低了使用门槛。",
        east="仿宋",
        size=10.5,
    )
    bold_p = doc.add_paragraph()
    bold_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bold_run = bold_p.add_run(
        "本段故意整段加粗，用来演示体检工具如何识别正文中的异常加粗，并给出可自动修复建议。"
    )
    bold_run.bold = True
    bold_run.font.size = Pt(10.5)
    bold_run.font.name = "Times New Roman"
    rpr = bold_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "仿宋")

    # 图与题注
    doc.add_picture(str(IMG), width=Cm(12))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "图1 系统总体架构", east="宋体", size=9, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 表与题注
    para(doc, "表1 实验数据对比", east="宋体", size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ("项目", "传统人工检查", "本文工具"),
        ("平均耗时", "35 分钟", "40 秒"),
        ("漏检项", "4.2 项", "0.6 项"),
        ("隐私风险", "需上传云端", "本地完成"),
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10.5)
                run.font.name = "Times New Roman"
                rp = run._element.get_or_add_rPr()
                rf = rp.find(qn("w:rFonts"))
                if rf is None:
                    rf = OxmlElement("w:rFonts")
                    rp.insert(0, rf)
                rf.set(qn("w:eastAsia"), "宋体")

    heading(doc, "2 系统设计", 1)
    para(
        doc,
        "系统由解析层、规则层、修复层和展示层组成。解析层读取文档的样式表与正文 XML，重建每个段落的有效格式；"
        "规则层把学校模板表达为可配置的 JSON 规则包；修复层只修改确定性强的问题，并保留其余内容不变。",
        east="仿宋",
        size=10.5,
    )
    # 故意跳级：一级标题直接到三级标题
    heading(doc, "2.1.1 数据库设计", 3)
    para(
        doc,
        "这里的标题层级故意从一级跳到三级，用来验证层级连续性检查。正确做法是先出现「2.1 模块划分」，再出现三级标题。",
        east="仿宋",
        size=10.5,
    )

    heading(doc, "3 结论", 1)
    para(
        doc,
        "论文格式体检把隐性的排版知识变成可执行的规则。下一步可以扩展学校规则包、增加更多自动修复项，"
        "并把检查能力接入写作流程，在提交前自动提醒。",
        east="仿宋",
        size=10.5,
    )

    # 谢辞（标题格式错误）
    para(doc, "谢  辞", east="宋体", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        doc,
        "感谢导师在论文写作与格式规范方面的耐心指导，也感谢同学在测试阶段提供的真实论文样本。",
        east="宋体",
        size=10.5,
    )

    # 参考文献（标题字体/对齐错误，条目含多种问题）
    para(doc, "参考文献", east="宋体", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    refs = [
        "[1] 张三. 论文格式自动检查研究. 天津大学学报, 2020, 53(3): 1-7.",
        "[2] 李四. 深度学习[M]. 北京: 科学出版社, 2019: 58-62.",
        "[3] 王五. 毕业论文排版方法[D]. 天津: 天津大学, 2021.",
        "[5] 赵六. 中文排版指南[M]. 北京: 人民邮电出版社, 2022.",
        "[6] Gedye R, Smith F, Westaway K, et al. Use of Microwave Ovens for Rapid Organic Synthesis[J]. Tetrahedron Lett, 1986, 27: 279.",
    ]
    for ref in refs:
        para(doc, ref, east="宋体", size=10.5)

    # 附录
    para(doc, "附  录", east="黑体", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        doc,
        "附录 A：规则包字段说明。规则包采用 JSON 描述页面、正文、标题、题注与参考文献的期望值，便于不同学院维护。",
        east="宋体",
        size=10.5,
    )

    doc.save(OUT)
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
